"""docx_to_pdf: failures are reported, and bad PDFs from an interrupted run are retried."""
import pytest

import Merge

GOOD_PDF = b"%PDF-1.4\n" + b"x" * 200


class FakeWord:
    """Stands in for Word.Application. Docs named in `fail` raise; `empty` ones save a 0-byte PDF."""
    def __init__(self, fail=(), empty=()):
        self.fail, self.empty = set(fail), set(empty)
        self.opened = []
        self.Visible = self.DisplayAlerts = None
        self.Documents = self

    def Open(self, path):
        name = path.replace("\\", "/").rsplit("/", 1)[-1]
        self.opened.append(name)
        if name in self.fail:
            raise RuntimeError("The RPC server is unavailable")
        return FakeDoc(name in self.empty)

    def Quit(self):
        pass


class FakeDoc:
    def __init__(self, empty):
        self.empty = empty

    def SaveAs(self, path, FileFormat):
        with open(path, "wb") as f:
            f.write(b"" if self.empty else GOOD_PDF)

    def Close(self, save):
        pass


@pytest.fixture
def word(monkeypatch):
    def install(**kw):
        fake = FakeWord(**kw)
        monkeypatch.setattr(Merge.win32com.client, "DispatchEx", lambda _name: fake)
        return fake
    return install


def make_docx(folder, *names):
    for n in names:
        (folder / f"{n}.docx").write_bytes(b"docx")


def test_failures_are_returned_and_bad_output_removed(tmp_path, word):
    make_docx(tmp_path, "1001", "1002", "1003")
    word(fail=["1002.docx"], empty=["1003.docx"])
    failures = Merge.docx_to_pdf(tmp_path, tmp_path)
    assert [name for name, _ in failures] == ["1002.docx", "1003.docx"]
    assert "RPC server" in failures[0][1]
    assert (tmp_path / "1001.pdf").exists()
    assert not (tmp_path / "1002.pdf").exists()
    assert not (tmp_path / "1003.pdf").exists()  # zero-byte output must not linger


def test_zero_byte_pdf_from_earlier_run_is_reconverted(tmp_path, word):
    make_docx(tmp_path, "2001")
    (tmp_path / "2001.pdf").write_bytes(b"")
    fake = word()
    assert Merge.docx_to_pdf(tmp_path, tmp_path) == []
    assert fake.opened == ["2001.docx"]
    assert Merge._validate_pdf(str(tmp_path / "2001.pdf"))[0]


def test_valid_pdf_is_not_reconverted(tmp_path, word):
    make_docx(tmp_path, "3001")
    (tmp_path / "3001.pdf").write_bytes(GOOD_PDF)
    fake = word()
    assert Merge.docx_to_pdf(tmp_path, tmp_path) == []
    assert fake.opened == []


def test_word_lock_files_are_ignored(tmp_path, word):
    make_docx(tmp_path, "4001", "~$4001")
    fake = word()
    Merge.docx_to_pdf(tmp_path, tmp_path)
    assert fake.opened == ["4001.docx"]


# ── createCoverLetter overwrite (the "Recreate All" button) ──────────────────

@pytest.fixture
def template(tmp_path):
    from docx import Document
    doc = Document()
    doc.add_paragraph("«Envelope_Name»")
    doc.add_paragraph("«Envelope_Number»")
    path = tmp_path / "template.docx"
    doc.save(path)
    return str(path)


def make_donor(name="New Name"):
    return Merge.Donor(pNum="1", pName="P", eNum="5", eName=name, aNum="9001")


def docx_text(path):
    from docx import Document
    return "\n".join(p.text for p in Document(str(path)).paragraphs)


def test_existing_pdf_is_kept_without_overwrite(tmp_path, template):
    out = tmp_path / "out"
    out.mkdir()
    (out / "9001.pdf").write_bytes(GOOD_PDF)
    Merge.createCoverLetter(make_donor(), str(out), template)
    assert not (out / "9001.docx").exists()


def test_overwrite_rebuilds_from_current_data(tmp_path, template):
    out = tmp_path / "out"
    out.mkdir()
    (out / "9001.pdf").write_bytes(GOOD_PDF)
    (out / "9001.docx").write_bytes(b"stale docx")
    donor = make_donor()
    donor.coverLetter = str(out / "9001.pdf")

    Merge.createCoverLetter(donor, str(out), template, overwrite=True)

    assert not (out / "9001.pdf").exists()  # old PDF gone, so docx_to_pdf will reconvert
    assert "New Name" in docx_text(out / "9001.docx")
    assert donor.coverLetter == ""


def test_overwrite_with_nothing_to_replace(tmp_path, template):
    out = tmp_path / "out"
    out.mkdir()
    Merge.createCoverLetter(make_donor(), str(out), template, overwrite=True)
    assert "New Name" in docx_text(out / "9001.docx")
